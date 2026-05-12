import json
import sqlite3
import io
import os
from datetime import datetime, date
from pathlib import Path
from flask import Flask, g, redirect, render_template, request, url_for, flash, session, jsonify
from flask import send_file
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "data" / "society.db"
SCHEMA_PATH = BASE_DIR / "schema.sql"
UPLOAD_DIR = BASE_DIR / "static" / "uploads" / "bills"
KYC_UPLOAD_DIR = BASE_DIR / "static" / "uploads" / "kyc"
DATABASE_URL = (os.environ.get("DATABASE_URL") or "").strip()
IS_POSTGRES = DATABASE_URL.startswith("postgresql://") or DATABASE_URL.startswith("postgres://")

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
ALLOWED_FLATS = [f"{floor}{unit:02d}" for floor in range(1, 10) for unit in range(1, 9)]
ALLOWED_BLOCKS = [f"Block-{i}" for i in range(1, 10)]

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "society-db-app-secret")
LOGIN_USERNAME = os.environ.get("LOGIN_USERNAME", "admin")
LOGIN_PASSWORD = os.environ.get("LOGIN_PASSWORD", "admin")
API_TOKEN_MAX_AGE_SECONDS = 60 * 60 * 24 * 7
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = (os.environ.get("SESSION_COOKIE_SECURE", "false").lower() == "true")


class DBRow(dict):
    def __init__(self, data: dict, values: list):
        super().__init__(data)
        self._values = values

    def __getitem__(self, key):
        if isinstance(key, int):
            return self._values[key]
        return super().__getitem__(key)


class DBResult:
    def __init__(self, rows: list[DBRow]):
        self._rows = rows

    def fetchall(self):
        return self._rows

    def fetchone(self):
        return self._rows[0] if self._rows else None


class DBConnection:
    def __init__(self, conn, backend: str):
        self._conn = conn
        self.backend = backend

    def _prepare_query(self, query: str):
        if self.backend == "postgres":
            return query.replace("?", "%s")
        return query

    def execute(self, query: str, params=None):
        params = params or []
        q = self._prepare_query(query)
        cur = self._conn.cursor()
        cur.execute(q, params)
        if not getattr(cur, "description", None):
            return DBResult([])
        columns = [d[0] for d in cur.description]
        rows = []
        for raw in cur.fetchall():
            if isinstance(raw, sqlite3.Row):
                values = [raw[c] for c in columns]
                data = {c: raw[c] for c in columns}
            else:
                values = list(raw)
                data = {c: raw[idx] for idx, c in enumerate(columns)}
            rows.append(DBRow(data, values))
        return DBResult(rows)

    def executescript(self, script: str):
        if self.backend == "sqlite":
            return self._conn.executescript(script)
        cur = self._conn.cursor()
        statements = [s.strip() for s in script.split(";") if s.strip()]
        for statement in statements:
            cur.execute(statement)
        return None

    def commit(self):
        self._conn.commit()

    def close(self):
        self._conn.close()


def _connect_db() -> DBConnection:
    if IS_POSTGRES:
        import psycopg
        raw = psycopg.connect(DATABASE_URL)
        return DBConnection(raw, "postgres")
    raw = sqlite3.connect(DB_PATH)
    raw.row_factory = sqlite3.Row
    return DBConnection(raw, "sqlite")


def get_db() -> DBConnection:
    if "db" not in g:
        g.db = _connect_db()
    return g.db


def is_logged_in() -> bool:
    return session.get("is_authenticated") is True


def _api_serializer() -> URLSafeTimedSerializer:
    return URLSafeTimedSerializer(app.secret_key, salt="society-mobile-api")


def _create_api_token(user_id: int, role: str, username: str) -> str:
    return _api_serializer().dumps({"user_id": int(user_id), "role": role, "username": username})


def _verify_api_token(token: str) -> dict | None:
    if not token:
        return None
    try:
        payload = _api_serializer().loads(token, max_age=API_TOKEN_MAX_AGE_SECONDS)
        if payload.get("username") == LOGIN_USERNAME or payload.get("user_id"):
            return payload
    except (BadSignature, SignatureExpired):
        return None
    return None


def _api_auth_payload() -> dict | None:
    auth_header = request.headers.get("Authorization", "")
    if auth_header.startswith("Bearer "):
        token = auth_header.split(" ", 1)[1].strip()
        return _verify_api_token(token)
    if is_logged_in():
        return {"username": LOGIN_USERNAME, "role": "ADMIN"}
    return None


def _get_current_user_row():
    payload = _api_auth_payload()
    if not payload:
        return None
    db = get_db()
    user_id = payload.get("user_id")
    if user_id:
        row = db.execute(
            """
            SELECT id, name, mobile, email, role, block, flat, status, password_hash,
                   vehicle_list, photo_url, living_from, rent_document_path, id_card_document_path
            FROM users WHERE id=?
            """,
            (user_id,),
        ).fetchone()
        if row:
            return row
    if payload.get("username") == LOGIN_USERNAME:
        return {
            "id": 1,
            "name": "Admin",
            "mobile": "",
            "email": "admin@societies.app",
            "role": "ADMIN",
            "block": "N/A",
            "flat": "N/A",
            "status": "APPROVED",
            "password_hash": "",
            "vehicle_list": "",
            "photo_url": "",
            "living_from": "",
            "rent_document_path": "",
            "id_card_document_path": "",
        }
    return None


def require_api_role(*roles):
    auth = require_api_auth()
    if auth:
        return auth
    user = _get_current_user_row()
    role = str((user or {}).get("role") or "").upper()
    if role not in {r.upper() for r in roles}:
        return jsonify({"ok": False, "error": "Forbidden"}), 403
    return None


def require_api_auth():
    if _api_auth_payload() is None:
        return jsonify({"ok": False, "error": "Unauthorized"}), 401
    return None


def require_login():
    if not is_logged_in():
        flash("Please login to access this page.", "error")
        return redirect(url_for("login", next=request.path))
    return None


@app.context_processor
def inject_auth_state():
    return {"is_authenticated": is_logged_in()}


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


@app.after_request
def add_cors_headers(response):
    origin = request.headers.get("Origin", "")
    if origin:
        response.headers["Access-Control-Allow-Origin"] = origin
        response.headers["Vary"] = "Origin"
    else:
        response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
    return response


@app.before_request
def handle_preflight():
    if request.method == "OPTIONS":
        return ("", 204)
    return None


def init_db() -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    KYC_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    if not IS_POSTGRES:
        DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    db = _connect_db()
    schema_file = BASE_DIR / ("schema_postgres.sql" if db.backend == "postgres" else "schema.sql")
    db.executescript(schema_file.read_text(encoding="utf-8"))
    if db.backend == "sqlite":
        cols = [r[1] for r in db.execute("PRAGMA table_info(payments)").fetchall()]
    else:
        cols = [r["column_name"] for r in db.execute(
            "SELECT column_name FROM information_schema.columns WHERE table_name='payments'"
        ).fetchall()]
    if "mode_of_payment" not in cols:
        db.execute("ALTER TABLE payments ADD COLUMN mode_of_payment TEXT")
    if db.backend == "postgres":
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS expenses (
              id BIGSERIAL PRIMARY KEY,
              transaction_date TEXT NOT NULL,
              item_name TEXT NOT NULL,
              quantity TEXT,
              amount DOUBLE PRECISION NOT NULL,
              payment_mode TEXT,
              paid_by TEXT,
              bill_image_path TEXT,
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
    else:
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS expenses (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              transaction_date TEXT NOT NULL,
              item_name TEXT NOT NULL,
              quantity TEXT,
              amount REAL NOT NULL,
              payment_mode TEXT,
              paid_by TEXT,
              bill_image_path TEXT,
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
    if db.backend == "sqlite":
        exp_cols = [r[1] for r in db.execute("PRAGMA table_info(expenses)").fetchall()]
    else:
        exp_cols = [r["column_name"] for r in db.execute(
            "SELECT column_name FROM information_schema.columns WHERE table_name='expenses'"
        ).fetchall()]
    if "quantity" not in exp_cols:
        db.execute("ALTER TABLE expenses ADD COLUMN quantity TEXT")

    if db.backend == "postgres":
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS owner_details (
              id BIGSERIAL PRIMARY KEY,
              property_id BIGINT NOT NULL UNIQUE REFERENCES properties(id) ON DELETE CASCADE,
              owner_name TEXT,
              owner_contact TEXT,
              is_occupied INTEGER NOT NULL DEFAULT 0,
              occupied_by TEXT NOT NULL DEFAULT 'OWNER',
              tenant_name TEXT,
              tenant_contact TEXT,
              tenant_vehicle_list TEXT,
              tenant_photo_url TEXT,
              tenant_living_from TEXT,
              tenant_guard_payment_details TEXT,
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
    else:
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS owner_details (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              property_id INTEGER NOT NULL UNIQUE,
              owner_name TEXT,
              owner_contact TEXT,
              is_occupied INTEGER NOT NULL DEFAULT 0,
              occupied_by TEXT NOT NULL DEFAULT 'OWNER',
              tenant_name TEXT,
              tenant_contact TEXT,
              tenant_vehicle_list TEXT,
              tenant_photo_url TEXT,
              tenant_living_from TEXT,
              tenant_guard_payment_details TEXT,
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              FOREIGN KEY(property_id) REFERENCES properties(id) ON DELETE CASCADE
            )
            """
        )

    db.execute(
        """
        INSERT INTO owner_details(
          property_id, owner_name, owner_contact, is_occupied, occupied_by, tenant_name, tenant_contact, tenant_vehicle_list, tenant_photo_url, tenant_living_from, tenant_guard_payment_details, updated_at
        )
        SELECT p.id, NULL, NULL, 0, 'OWNER', NULL, NULL, NULL, NULL, NULL, NULL, CURRENT_TIMESTAMP
        FROM properties p
        WHERE NOT EXISTS (SELECT 1 FROM owner_details od WHERE od.property_id = p.id)
        """
    )

    if db.backend == "postgres":
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
              id BIGSERIAL PRIMARY KEY,
              name TEXT NOT NULL,
              mobile TEXT NOT NULL UNIQUE,
              email TEXT NOT NULL UNIQUE,
              password_hash TEXT NOT NULL,
              role TEXT NOT NULL DEFAULT 'TENANT',
              block TEXT NOT NULL,
              flat TEXT NOT NULL,
              status TEXT NOT NULL DEFAULT 'APPROVED',
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS registration_requests (
              id BIGSERIAL PRIMARY KEY,
              name TEXT NOT NULL,
              mobile TEXT NOT NULL UNIQUE,
              email TEXT NOT NULL UNIQUE,
              password_hash TEXT NOT NULL,
              block TEXT NOT NULL,
              flat TEXT NOT NULL,
              status TEXT NOT NULL DEFAULT 'PENDING',
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
    else:
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              name TEXT NOT NULL,
              mobile TEXT NOT NULL UNIQUE,
              email TEXT NOT NULL UNIQUE,
              password_hash TEXT NOT NULL,
              role TEXT NOT NULL DEFAULT 'TENANT',
              block TEXT NOT NULL,
              flat TEXT NOT NULL,
              status TEXT NOT NULL DEFAULT 'APPROVED',
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS registration_requests (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              name TEXT NOT NULL,
              mobile TEXT NOT NULL UNIQUE,
              email TEXT NOT NULL UNIQUE,
              password_hash TEXT NOT NULL,
              block TEXT NOT NULL,
              flat TEXT NOT NULL,
              status TEXT NOT NULL DEFAULT 'PENDING',
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )

    admin_exists = db.execute("SELECT id FROM users WHERE LOWER(email)=LOWER(?)", ("admin@societies.app",)).fetchone()
    if not admin_exists:
        db.execute(
            """
            INSERT INTO users(name, mobile, email, password_hash, role, block, flat, status, updated_at)
            VALUES (?, ?, ?, ?, 'ADMIN', 'N/A', 'N/A', 'APPROVED', CURRENT_TIMESTAMP)
            """,
            ("Admin", "0000000000", "admin@societies.app", generate_password_hash(LOGIN_PASSWORD)),
        )

    # Backward-compatible column migrations for user profile and registration docs.
    migration_map = {
        "users": ["vehicle_list", "photo_url", "living_from", "rent_document_path", "id_card_document_path"],
        "registration_requests": ["living_from", "rent_document_path", "id_card_document_path"],
    }
    for table, cols in migration_map.items():
        if db.backend == "sqlite":
            existing_cols = [r[1] for r in db.execute(f"PRAGMA table_info({table})").fetchall()]
        else:
            existing_cols = [r["column_name"] for r in db.execute(
                "SELECT column_name FROM information_schema.columns WHERE table_name=?",
                (table,),
            ).fetchall()]
        for col in cols:
            if col not in existing_cols:
                db.execute(f"ALTER TABLE {table} ADD COLUMN {col} TEXT")
    db.commit()
    db.close()


def _save_uploaded_kyc(file_storage):
    if not file_storage or not file_storage.filename:
        return None, None
    safe = secure_filename(file_storage.filename)
    ext = Path(safe).suffix.lower()
    if ext not in {".png", ".jpg", ".jpeg", ".webp", ".pdf"}:
        return None, "Document must be image or PDF."
    unique_name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{safe}"
    out_path = KYC_UPLOAD_DIR / unique_name
    out_path.parent.mkdir(parents=True, exist_ok=True)
    file_storage.save(out_path)
    return f"uploads/kyc/{unique_name}", None


def _find_resident_directory_match(db: sqlite3.Connection, block: str, flat: str):
    rows = db.execute("SELECT raw_json FROM resident_directory ORDER BY id DESC").fetchall()
    block_norm = str(block or "").strip().upper()
    flat_norm = str(flat or "").strip().upper()
    flat_digits = "".join(ch for ch in flat_norm if ch.isdigit())
    for r in rows:
        try:
            record = json.loads(r["raw_json"])
        except Exception:
            continue
        values = {str(k).strip().upper(): str(v).strip() for k, v in record.items()}
        joined = " | ".join(str(v).upper() for v in values.values())
        if block_norm and block_norm not in joined:
            continue
        if flat_norm and flat_norm in joined:
            return record
        if flat_digits and flat_digits in "".join(ch for ch in joined if ch.isdigit()):
            return record
    return None


def _normalize_date_to_iso(date_str: str | None) -> str | None:
    if not date_str:
        return None
    val = date_str.strip()
    if not val:
        return None
    # already YYYY-MM-DD
    if len(val) == 10 and val[4] == "-" and val[7] == "-":
        return val
    # MM/DD/YYYY
    if "/" in val:
        parts = val.split("/")
        if len(parts) == 3:
            mm, dd, yyyy = parts
            try:
                return f"{int(yyyy):04d}-{int(mm):02d}-{int(dd):02d}"
            except Exception:
                return None
    return None


def _normalize_date_for_storage(date_str: str | None) -> str | None:
    # store as MM/DD/YYYY in DB to stay compatible with imported data
    iso = _normalize_date_to_iso(date_str)
    if not iso:
        return None
    yyyy, mm, dd = iso.split("-")
    return f"{mm}/{dd}/{yyyy}"


def _format_db_date_to_iso(date_str: str | None) -> str:
    iso = _normalize_date_to_iso(date_str)
    return iso or ""


def _parse_db_date(date_str: str | None) -> date | None:
    iso = _normalize_date_to_iso(date_str)
    if not iso:
        return None
    try:
        return datetime.strptime(iso, "%Y-%m-%d").date()
    except Exception:
        return None


def _save_uploaded_bill(file_storage):
    if not file_storage or not file_storage.filename:
        return None, None
    safe = secure_filename(file_storage.filename)
    ext = Path(safe).suffix.lower()
    if ext not in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".pdf"}:
        return None, "Bill must be an image or PDF file."
    unique_name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{safe}"
    out_path = UPLOAD_DIR / unique_name
    out_path.parent.mkdir(parents=True, exist_ok=True)
    file_storage.save(out_path)
    return f"uploads/bills/{unique_name}", None


def _build_public_summary(db: sqlite3.Connection):
    today = datetime.now().date()
    current_year = today.year
    current_month = today.month

    done_rows = db.execute(
        """
        SELECT p.block, p.flat, pay.year, pay.month, pay.amount, pay.payment_date, pay.mode_of_payment
        FROM payments pay
        JOIN properties p ON p.id = pay.property_id
        WHERE UPPER(COALESCE(pay.status, 'PENDING')) = 'DONE'
        """
    ).fetchall()

    month_count = 0
    month_amount = 0.0
    today_count = 0
    today_amount = 0.0
    year_count = 0
    year_amount = 0.0
    total_collection = 0.0
    block_month = {}
    all_years = {}
    mode_split_month = {"ONLINE": 0.0, "CASH": 0.0, "OTHER": 0.0}
    recent_payments = []
    trend_map = {}

    for r in done_rows:
        y = int(r["year"] or 0)
        m = int(r["month"] or 0)
        amt = float(r["amount"] or 0)
        block = r["block"] or "-"
        payment_dt = _parse_db_date(r["payment_date"])
        mode = (r["mode_of_payment"] or "").strip().upper()

        if y not in trend_map:
            trend_map[y] = [0.0] * 12
        if 1 <= m <= 12:
            trend_map[y][m - 1] += amt

        if y == current_year:
            year_count += 1
            year_amount += amt

        if y == current_year and m == current_month:
            month_count += 1
            month_amount += amt
            block_month[block] = block_month.get(block, {"count": 0, "amount": 0.0})
            block_month[block]["count"] += 1
            block_month[block]["amount"] += amt
            if mode == "ONLINE":
                mode_split_month["ONLINE"] += amt
            elif mode == "CASH":
                mode_split_month["CASH"] += amt
            else:
                mode_split_month["OTHER"] += amt

        if payment_dt and payment_dt == today:
            today_count += 1
            today_amount += amt

        all_years[y] = all_years.get(y, {"count": 0, "amount": 0.0})
        all_years[y]["count"] += 1
        all_years[y]["amount"] += amt
        total_collection += amt

        recent_payments.append({
            "block": block,
            "flat": r["flat"],
            "amount": amt,
            "date": r["payment_date"] or "-",
            "mode": mode or "-",
        })

    recent_payments = recent_payments[:10]
    block_month_list = sorted(block_month.items(), key=lambda item: item[0])
    all_years_list = sorted(all_years.items(), key=lambda item: item[0], reverse=True)

    total_flats = int(db.execute("SELECT COUNT(*) FROM properties").fetchone()[0])
    pending_this_month_flats = max(total_flats - int(month_count), 0)

    # Top pending blocks
    block_total_rows = db.execute("SELECT block, COUNT(*) AS flat_count FROM properties GROUP BY block").fetchall()
    block_total_flats = {str(r["block"]): int(r["flat_count"] or 0) for r in block_total_rows}
    block_paid_month = {b: info["count"] for b, info in block_month_list}
    pending_blocks = []
    for block, total in block_total_flats.items():
        paid = int(block_paid_month.get(block, 0))
        pending = max(int(total) - paid, 0)
        completion = (paid / total * 100.0) if total else 0.0
        pending_blocks.append({
            "block": block,
            "total_flats": int(total),
            "paid_flats": paid,
            "pending_flats": pending,
            "completion_pct": completion,
        })
    pending_blocks.sort(key=lambda x: (x["completion_pct"], -x["pending_flats"], x["block"]))
    top_pending_blocks = pending_blocks[:5]

    trend_years = sorted(trend_map.keys(), reverse=False)
    trend_datasets = [
        {"year": y, "values": trend_map[y]} for y in trend_years
    ]

    return {
        "today": today.strftime("%d %b %Y"),
        "month_name": MONTHS[current_month - 1],
        "year": current_year,
        "month_count": month_count,
        "month_amount": month_amount,
        "today_count": today_count,
        "today_amount": today_amount,
        "year_count": year_count,
        "year_amount": year_amount,
        "total_collection": total_collection,
        "block_month": block_month_list,
        "all_years": all_years_list,
        "mode_split_month": mode_split_month,
        "total_flats": total_flats,
        "pending_this_month_flats": pending_this_month_flats,
        "recent_payments": recent_payments,
        "top_pending_blocks": top_pending_blocks,
        "trend_datasets": trend_datasets,
    }


def _extract_profile_from_directory(record: dict):
    if not record:
        return {}
    normalized = {str(k).strip().lower(): str(v).strip() for k, v in record.items()}

    def pick(*keys):
        for k in normalized:
            for needle in keys:
                if needle in k and normalized[k]:
                    return normalized[k]
        return ""

    return {
        "occupant_name": pick("occupant", "resident", "name"),
        "occupant_phone": pick("occupant phone", "phone", "mobile", "contact"),
        "owner_name": pick("owner name", "owner"),
        "owner_phone": pick("owner phone", "owner mobile", "owner contact"),
        "tenant_name": pick("tenant name", "tenant"),
        "tenant_phone": pick("tenant phone", "tenant mobile", "tenant contact"),
        "vehicle_number": pick("vehicle", "car", "bike", "registration"),
        "notes": "",
    }


@app.route("/")
def home():
    return redirect(url_for("public_dashboard"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        password = request.form.get("password") or ""
        next_url = request.form.get("next") or url_for("public_dashboard")
        if username == LOGIN_USERNAME and password == LOGIN_PASSWORD:
            session["is_authenticated"] = True
            flash("Login successful.", "success")
            return redirect(next_url)
        flash("Invalid username or password.", "error")
    return render_template("login.html", active_page="login", next_url=request.args.get("next", ""))


@app.route("/logout")
def logout():
    session.pop("is_authenticated", None)
    flash("Logged out successfully.", "success")
    return redirect(url_for("public_dashboard"))


@app.route("/public-dashboard")
def public_dashboard():
    db = get_db()
    summary = _build_public_summary(db)

    notices = [
        "Monthly maintenance due date is 10th of every month.",
        "Please share payment screenshot/transaction reference after online transfer.",
    ]
    important_contacts = [
        {"label": "Society Office", "value": "+91-90000-00001"},
        {"label": "Security Supervisor", "value": "+91-90000-00002"},
        {"label": "Electrician", "value": "+91-90000-00003"},
    ]

    return render_template(
        "public_dashboard.html",
        active_page="public_dashboard",
        today=summary["today"],
        month_name=summary["month_name"],
        year=summary["year"],
        month_count=summary["month_count"],
        month_amount=summary["month_amount"],
        today_count=summary["today_count"],
        today_amount=summary["today_amount"],
        year_count=summary["year_count"],
        year_amount=summary["year_amount"],
        block_month=summary["block_month"],
        all_years=summary["all_years"],
        notices=notices,
        contacts=important_contacts,
        mode_split_month=summary["mode_split_month"],
        total_flats=summary["total_flats"],
        pending_this_month_flats=summary["pending_this_month_flats"],
        recent_payments=summary["recent_payments"],
        top_pending_blocks=summary["top_pending_blocks"],
        trend_datasets=summary["trend_datasets"],
        month_labels=MONTHS,
    )


@app.route("/public-dashboard/export/excel")
def public_dashboard_export_excel():
    db = get_db()
    summary = _build_public_summary(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    ws.append(["Date", summary["today"]])
    ws.append(["Month", f'{summary["month_name"]} {summary["year"]}'])
    ws.append(["Payments This Month", summary["month_count"]])
    ws.append(["Collection This Month", summary["month_amount"]])
    ws.append(["Payments Today", summary["today_count"]])
    ws.append(["Collection Today", summary["today_amount"]])
    ws.append(["Payments This Year", summary["year_count"]])
    ws.append(["Collection This Year", summary["year_amount"]])

    ws2 = wb.create_sheet("Block-Month")
    ws2.append(["Block", "Payments", "Collection"])
    for block, info in summary["block_month"]:
        ws2.append([block, info["count"], info["amount"]])

    ws3 = wb.create_sheet("Yearly")
    ws3.append(["Year", "Payments", "Collection"])
    for y, info in summary["all_years"]:
        ws3.append([y, info["count"], info["amount"]])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"society_summary_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/public-dashboard/export/pdf")
def public_dashboard_export_pdf():
    db = get_db()
    summary = _build_public_summary(db)

    output = io.BytesIO()
    c = canvas.Canvas(output, pagesize=A4)
    width, height = A4
    y = height - 40
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, "MIG-1 Society Public Summary")
    y -= 24
    c.setFont("Helvetica", 11)
    lines = [
        f"Date: {summary['today']}",
        f"Month: {summary['month_name']} {summary['year']}",
        f"Payments This Month: {summary['month_count']}",
        f"Collection This Month: Rs {summary['month_amount']:.0f}",
        f"Payments Today: {summary['today_count']}",
        f"Collection Today: Rs {summary['today_amount']:.0f}",
        f"Payments This Year: {summary['year_count']}",
        f"Collection This Year: Rs {summary['year_amount']:.0f}",
        "",
        "Top Pending Blocks:",
    ]
    for line in lines:
        c.drawString(40, y, line)
        y -= 16
    for b in summary["top_pending_blocks"]:
        c.drawString(50, y, f"{b['block']} | Paid: {b['paid_flats']} / {b['total_flats']} | Pending: {b['pending_flats']} | Completion: {b['completion_pct']:.1f}%")
        y -= 14
        if y < 60:
            c.showPage()
            y = height - 40
    c.save()
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"society_summary_{datetime.now().strftime('%Y%m%d')}.pdf",
        mimetype="application/pdf",
    )


@app.route("/payments")
def payments():
    db = get_db()

    year = int(request.args.get("year", 2026))
    month = int(request.args.get("month", datetime.now().month))
    block = request.args.get("block", "ALL")
    search = request.args.get("search", "").strip()
    payment_date_filter = request.args.get("payment_date", "").strip()

    years = [row[0] for row in db.execute("SELECT DISTINCT year FROM payments ORDER BY year DESC").fetchall()]
    if not years:
        years = [2026, 2025]

    blocks = [row[0] for row in db.execute("SELECT DISTINCT block FROM properties ORDER BY block").fetchall()]

    query = """
    SELECT p.id AS property_id, p.block, p.flat, pay.amount, pay.status, pay.year, pay.month, pay.payment_date
    FROM properties p
    LEFT JOIN payments pay ON pay.property_id = p.id AND pay.year = ?
    WHERE 1=1
    """
    params = [year]

    if block != "ALL":
        query += " AND p.block = ?"
        params.append(block)

    if search:
        query += " AND (p.flat LIKE ? OR p.block LIKE ?)"
        like_q = f"%{search}%"
        params.extend([like_q, like_q])
    if payment_date_filter:
        db_date_filter = _normalize_date_for_storage(payment_date_filter)
        if db_date_filter:
            query += " AND pay.payment_date = ?"
            params.append(db_date_filter)

    query += " ORDER BY p.block, CAST(p.flat AS INTEGER), pay.month"

    raw_rows = db.execute(query, params).fetchall()

    flats_map = {}
    for r in raw_rows:
        pid = r["property_id"]
        if pid not in flats_map:
            flats_map[pid] = {
                "property_id": pid,
                "block": r["block"],
                "flat": r["flat"],
                "months": [None for _ in range(12)],
                "total": 0.0,
            }

        m = r["month"]
        if not m or m < 1 or m > 12:
            continue
        amount = float(r["amount"] or 0)
        status = (r["status"] or "PENDING").upper()
        payment_date = r["payment_date"]

        flats_map[pid]["months"][m - 1] = {
            "amount": amount,
            "status": status,
            "payment_date": payment_date,
        }
        if status == "DONE":
            flats_map[pid]["total"] += amount

    table_rows = sorted(flats_map.values(), key=lambda x: (x["block"], str(x["flat"])))

    monthly_totals = [0.0 for _ in range(12)]
    paid_count = 0
    locked_count = 0
    for row in table_rows:
        selected_month_data = row["months"][month - 1]
        if selected_month_data:
            if selected_month_data["status"] == "DONE":
                paid_count += 1
            if selected_month_data["status"] == "LOCKED":
                locked_count += 1

        for i, month_data in enumerate(row["months"]):
            if month_data and month_data["status"] == "DONE":
                monthly_totals[i] += float(month_data["amount"] or 0)

    total_amount = monthly_totals[month - 1] if 1 <= month <= 12 else 0.0
    total_flats = len(table_rows)
    unpaid_count = max(total_flats - paid_count - locked_count, 0)

    return render_template(
        "payments.html",
        rows=table_rows,
        years=years,
        blocks=blocks,
        filters={"year": year, "month": month, "block": block, "search": search},
        payment_date_filter=payment_date_filter,
        months=MONTHS,
        total_amount=total_amount,
        paid_count=paid_count,
        unpaid_count=unpaid_count,
        locked_count=locked_count,
        total_flats=total_flats,
        monthly_totals=monthly_totals,
        active_page="payments",
    )


@app.route("/payment/<int:property_id>/<int:year>/<int:month>/edit", methods=["GET", "POST"])
def edit_payment(property_id: int, year: int, month: int):
    auth_redirect = require_login()
    if auth_redirect:
        return auth_redirect
    db = get_db()
    return_to = request.values.get("return_to", url_for("payments"))

    if month < 1 or month > 12:
        flash("Invalid month selected.", "error")
        return redirect(return_to)

    property_row = db.execute(
        "SELECT id, block, flat FROM properties WHERE id = ?",
        (property_id,),
    ).fetchone()
    if not property_row:
        flash("Flat record not found.", "error")
        return redirect(return_to)

    if request.method == "POST":
        action = (request.form.get("action") or "save").strip().lower()
        if action == "delete":
            db.execute(
                "DELETE FROM payments WHERE property_id=? AND year=? AND month=?",
                (property_id, year, month),
            )
            db.commit()
            flash("Payment deleted successfully.", "success")
            return redirect(return_to)

        flat = request.form.get("flat", "").strip()
        amount = float(request.form.get("amount", "0") or 0)
        payment_date = _normalize_date_for_storage(request.form.get("payment_date", "").strip() or None)
        mode_of_payment = request.form.get("mode_of_payment", "").strip() or None
        received_by = request.form.get("received_by", "").strip() or None
        status = request.form.get("status", "PENDING").strip().upper() or "PENDING"
        notes = request.form.get("notes", "").strip() or None

        if mode_of_payment:
            mode_of_payment = mode_of_payment.upper()
        if mode_of_payment == "ONLINE":
            received_by = None
        if mode_of_payment == "CASH" and not received_by:
            flash("Please enter receiver name for cash payment.", "error")
            return redirect(return_to)
        if mode_of_payment not in {"CASH", "ONLINE", None}:
            flash("Mode of payment must be CASH or ONLINE.", "error")
            return redirect(return_to)

        if flat not in ALLOWED_FLATS:
            flash("Invalid flat selected. Please choose from the provided flat list.", "error")
            return redirect(return_to)

        target_property_id = property_id
        if flat != property_row["flat"]:
            existing = db.execute(
                "SELECT id FROM properties WHERE block=? AND flat=?",
                (property_row["block"], flat),
            ).fetchone()
            if existing:
                target_property_id = int(existing["id"])
            else:
                db.execute(
                    "UPDATE properties SET flat=?, updated_at=CURRENT_TIMESTAMP WHERE id=?",
                    (flat, property_id),
                )
                target_property_id = property_id

        db.execute(
            """
            INSERT INTO payments(property_id, year, month, amount, payment_date, status, notes, source, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, 'manual', CURRENT_TIMESTAMP)
            ON CONFLICT(property_id, year, month) DO UPDATE SET
              amount=excluded.amount,
              payment_date=excluded.payment_date,
              status=excluded.status,
              notes=excluded.notes,
              mode_of_payment=excluded.mode_of_payment,
              source='manual',
              updated_at=CURRENT_TIMESTAMP
            """,
            (target_property_id, year, month, amount if status == "DONE" else 0, payment_date, status, notes),
        )
        db.execute(
            "UPDATE payments SET mode_of_payment=?, notes=COALESCE(?, notes) WHERE property_id=? AND year=? AND month=?",
            (mode_of_payment, notes, target_property_id, year, month),
        )
        if received_by:
            db.execute(
                "UPDATE payments SET notes=TRIM(COALESCE(notes,'') || CASE WHEN COALESCE(notes,'')='' THEN '' ELSE ' | ' END || ?) WHERE property_id=? AND year=? AND month=?",
                (f"Received By: {received_by}", target_property_id, year, month),
            )
        if target_property_id != property_id:
            db.execute(
                "DELETE FROM payments WHERE property_id=? AND year=? AND month=?",
                (property_id, year, month),
            )
        db.commit()
        flash("Payment details updated successfully.", "success")
        return redirect(return_to)

    payment_row = db.execute(
        """
        SELECT amount, payment_date, status, notes, mode_of_payment
        FROM payments
        WHERE property_id = ? AND year = ? AND month = ?
        """,
        (property_id, year, month),
    ).fetchone()

    return render_template(
        "payment_edit.html",
        active_page="entry",
        row=property_row,
        payment=payment_row,
        year=year,
        month=month,
        months=MONTHS,
        allowed_flats=ALLOWED_FLATS,
        payment_date_iso=_format_db_date_to_iso(payment_row["payment_date"]) if payment_row else "",
        received_by=(
            (payment_row["notes"].split("Received By: ")[-1].strip() if payment_row and payment_row["notes"] and "Received By:" in payment_row["notes"] else "")
        ),
        return_to=return_to,
    )


@app.route("/entry", methods=["GET", "POST"])
def entry():
    auth_redirect = require_login()
    if auth_redirect:
        return auth_redirect
    db = get_db()
    search_text = request.args.get("q", "").strip()
    date_filter = request.args.get("payment_date", "").strip()

    blocks = [row[0] for row in db.execute("SELECT DISTINCT block FROM properties ORDER BY block").fetchall()]
    years = [row[0] for row in db.execute("SELECT DISTINCT year FROM payments ORDER BY year DESC").fetchall()]
    if not years:
        years = [2026, 2025]

    query = """
    SELECT p.id AS property_id, p.block, p.flat, pay.year, pay.month, pay.amount, pay.payment_date, pay.source, pay.mode_of_payment, pay.status
    FROM payments pay
    JOIN properties p ON p.id = pay.property_id
    WHERE (pay.status IS NULL OR UPPER(pay.status)='DONE')
    """
    params = []
    if search_text:
        query += " AND (p.block LIKE ? OR p.flat LIKE ? OR pay.mode_of_payment LIKE ?)"
        like_q = f"%{search_text}%"
        params.extend([like_q, like_q, like_q])
    if date_filter:
        db_date_filter = _normalize_date_for_storage(date_filter)
        if db_date_filter:
            query += " AND pay.payment_date = ?"
            params.append(db_date_filter)
    query += " ORDER BY pay.updated_at DESC LIMIT 50"
    recent = db.execute(query, params).fetchall()

    return render_template(
        "entry.html",
        blocks=blocks,
        years=years,
        months=MONTHS,
        recent=recent,
        allowed_flats=ALLOWED_FLATS,
        search_text=search_text,
        date_filter=date_filter,
        active_page="entry",
    )


@app.route("/entry/new", methods=["GET", "POST"])
def entry_new():
    auth_redirect = require_login()
    if auth_redirect:
        return auth_redirect
    db = get_db()
    if request.method == "POST":
        block = request.form.get("block", "").strip()
        flat = request.form.get("flat", "").strip()
        year = int(request.form.get("year", "0"))
        month = int(request.form.get("month", "0"))
        amount = float(request.form.get("amount", "0") or 0)
        payment_date = _normalize_date_for_storage(request.form.get("payment_date", "").strip() or None)
        status = "DONE"
        notes = request.form.get("notes", "").strip() or None
        mode_of_payment = request.form.get("mode_of_payment", "").strip() or None
        received_by = request.form.get("received_by", "").strip() or None

        if mode_of_payment:
            mode_of_payment = mode_of_payment.upper()
        if mode_of_payment == "ONLINE":
            received_by = None
        if mode_of_payment == "CASH" and not received_by:
            flash("Please enter receiver name for cash payment.", "error")
            return redirect(url_for("entry_new"))
        if mode_of_payment not in {"CASH", "ONLINE", None}:
            flash("Mode of payment must be CASH or ONLINE.", "error")
            return redirect(url_for("entry_new"))

        if not block or not flat or not (1 <= month <= 12) or year < 2000:
            flash("Please provide valid block, flat, year and month.", "error")
            return redirect(url_for("entry_new"))
        if flat not in ALLOWED_FLATS:
            flash("Invalid flat selected. Please choose from the provided flat list.", "error")
            return redirect(url_for("entry_new"))

        db.execute(
            """
            INSERT INTO properties(block, flat, updated_at)
            VALUES (?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(block, flat) DO UPDATE SET updated_at=CURRENT_TIMESTAMP
            """,
            (block, flat),
        )
        property_id = db.execute(
            "SELECT id FROM properties WHERE block=? AND flat=?",
            (block, flat),
        ).fetchone()[0]

        db.execute(
            """
            INSERT INTO payments(property_id, year, month, amount, payment_date, status, notes, source, mode_of_payment, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, 'manual', ?, CURRENT_TIMESTAMP)
            ON CONFLICT(property_id, year, month) DO UPDATE SET
              amount=excluded.amount,
              payment_date=excluded.payment_date,
              status=excluded.status,
              notes=excluded.notes,
              mode_of_payment=excluded.mode_of_payment,
              source='manual',
              updated_at=CURRENT_TIMESTAMP
            """,
            (property_id, year, month, amount, payment_date, status, notes, mode_of_payment),
        )
        if received_by:
            db.execute(
                "UPDATE payments SET notes=TRIM(COALESCE(notes,'') || CASE WHEN COALESCE(notes,'')='' THEN '' ELSE ' | ' END || ?) WHERE property_id=? AND year=? AND month=?",
                (f"Received By: {received_by}", property_id, year, month),
            )
        db.commit()
        flash("Payment saved successfully.", "success")
        return redirect(url_for("entry"))

    blocks = [row[0] for row in db.execute("SELECT DISTINCT block FROM properties ORDER BY block").fetchall()]
    years = [row[0] for row in db.execute("SELECT DISTINCT year FROM payments ORDER BY year DESC").fetchall()]
    if not years:
        years = [2026, 2025]
    return render_template(
        "entry_new.html",
        blocks=blocks,
        years=years,
        months=MONTHS,
        allowed_flats=ALLOWED_FLATS,
        today_iso=datetime.now().strftime("%Y-%m-%d"),
        active_page="entry",
    )


@app.route("/flat/<int:property_id>", methods=["GET", "POST"])
def flat_details(property_id: int):
    auth_redirect = require_login()
    if auth_redirect:
        return auth_redirect
    db = get_db()
    property_row = db.execute(
        "SELECT id, block, flat FROM properties WHERE id=?",
        (property_id,),
    ).fetchone()
    if not property_row:
        flash("Flat record not found.", "error")
        return redirect(url_for("payments"))

    if request.method == "POST":
        db.execute(
            """
            INSERT INTO occupant_profiles(
              property_id, occupant_name, occupant_phone, owner_name, owner_phone,
              tenant_name, tenant_phone, vehicle_number, notes, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(property_id) DO UPDATE SET
              occupant_name=excluded.occupant_name,
              occupant_phone=excluded.occupant_phone,
              owner_name=excluded.owner_name,
              owner_phone=excluded.owner_phone,
              tenant_name=excluded.tenant_name,
              tenant_phone=excluded.tenant_phone,
              vehicle_number=excluded.vehicle_number,
              notes=excluded.notes,
              updated_at=CURRENT_TIMESTAMP
            """,
            (
                property_id,
                request.form.get("occupant_name", "").strip() or None,
                request.form.get("occupant_phone", "").strip() or None,
                request.form.get("owner_name", "").strip() or None,
                request.form.get("owner_phone", "").strip() or None,
                request.form.get("tenant_name", "").strip() or None,
                request.form.get("tenant_phone", "").strip() or None,
                request.form.get("vehicle_number", "").strip() or None,
                request.form.get("notes", "").strip() or None,
            ),
        )
        db.commit()
        flash("Occupant details updated successfully.", "success")
        return redirect(url_for("flat_details", property_id=property_id))

    profile = db.execute(
        "SELECT * FROM occupant_profiles WHERE property_id=?",
        (property_id,),
    ).fetchone()

    if not profile:
        derived = _extract_profile_from_directory(_find_resident_directory_match(db, property_row["block"], property_row["flat"]))
        if any(derived.values()):
            db.execute(
                """
                INSERT INTO occupant_profiles(
                  property_id, occupant_name, occupant_phone, owner_name, owner_phone,
                  tenant_name, tenant_phone, vehicle_number, notes, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(property_id) DO UPDATE SET
                  occupant_name=excluded.occupant_name,
                  occupant_phone=excluded.occupant_phone,
                  owner_name=excluded.owner_name,
                  owner_phone=excluded.owner_phone,
                  tenant_name=excluded.tenant_name,
                  tenant_phone=excluded.tenant_phone,
                  vehicle_number=excluded.vehicle_number,
                  notes=excluded.notes,
                  updated_at=CURRENT_TIMESTAMP
                """,
                (
                    property_id,
                    derived.get("occupant_name") or None,
                    derived.get("occupant_phone") or None,
                    derived.get("owner_name") or None,
                    derived.get("owner_phone") or None,
                    derived.get("tenant_name") or None,
                    derived.get("tenant_phone") or None,
                    derived.get("vehicle_number") or None,
                    derived.get("notes") or None,
                ),
            )
            db.commit()
            profile = db.execute(
                "SELECT * FROM occupant_profiles WHERE property_id=?",
                (property_id,),
            ).fetchone()

    payment_history = db.execute(
        """
        SELECT id, year, month, amount, status, payment_date, notes, source, mode_of_payment
        FROM payments
        WHERE property_id=? AND UPPER(COALESCE(status,'PENDING'))='DONE'
        ORDER BY year DESC, month DESC
        """,
        (property_id,),
    ).fetchall()

    directory_match = _find_resident_directory_match(db, property_row["block"], property_row["flat"])
    return render_template(
        "flat_details.html",
        active_page="payments",
        row=property_row,
        profile=profile,
        directory_match=directory_match,
        payment_history=payment_history,
        months=MONTHS,
    )


@app.route("/directory")
def directory():
    auth_redirect = require_login()
    if auth_redirect:
        return auth_redirect
    db = get_db()
    search = request.args.get("search", "").strip().lower()
    rows = db.execute("SELECT id, raw_json FROM resident_directory ORDER BY id DESC").fetchall()

    data = []
    for r in rows:
        record = json.loads(r["raw_json"])
        record["_id"] = r["id"]
        data.append(record)
    if search:
        def match(record):
            txt = " ".join(str(v) for v in record.values()).lower()
            return search in txt
        data = [d for d in data if match(d)]

    headers = [h for h in (list(data[0].keys()) if data else []) if h != "_id"]
    return render_template(
        "directory.html",
        headers=headers,
        data=data,
        search=search,
        active_page="directory",
    )


@app.route("/directory/<int:row_id>/edit", methods=["GET", "POST"])
def directory_edit(row_id: int):
    auth_redirect = require_login()
    if auth_redirect:
        return auth_redirect
    db = get_db()
    row = db.execute("SELECT id, raw_json FROM resident_directory WHERE id=?", (row_id,)).fetchone()
    if not row:
        flash("Directory record not found.", "error")
        return redirect(url_for("directory"))
    record = json.loads(row["raw_json"])

    if request.method == "POST":
        updated = {}
        for k in record.keys():
            updated[k] = request.form.get(k, "").strip()
        db.execute(
            "UPDATE resident_directory SET raw_json=?, updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (json.dumps(updated, ensure_ascii=True), row_id),
        )
        db.commit()
        flash("Directory record updated successfully.", "success")
        return redirect(url_for("directory"))

    return render_template(
        "directory_edit.html",
        active_page="directory",
        row_id=row_id,
        record=record,
    )


@app.route("/expenses")
def expenses():
    db = get_db()
    search_item = request.args.get("item", "").strip()
    search_date = request.args.get("transaction_date", "").strip()
    sort = (request.args.get("sort", "desc") or "desc").lower()
    page = max(int(request.args.get("page", "1") or "1"), 1)
    per_page = 10

    order = "DESC" if sort != "asc" else "ASC"
    where = " WHERE 1=1 "
    params = []
    if search_item:
        where += " AND item_name LIKE ? "
        params.append(f"%{search_item}%")
    if search_date:
        db_date = _normalize_date_for_storage(search_date)
        if db_date:
            where += " AND transaction_date = ? "
            params.append(db_date)

    total_rows = db.execute(f"SELECT COUNT(*) FROM expenses {where}", params).fetchone()[0]
    total_pages = max((total_rows + per_page - 1) // per_page, 1)
    if page > total_pages:
        page = total_pages
    offset = (page - 1) * per_page

    rows = db.execute(
        f"""
        SELECT id, transaction_date, item_name, amount, payment_mode, paid_by, bill_image_path, created_at
        FROM expenses
        {where}
        ORDER BY transaction_date {order}, id {order}
        LIMIT ? OFFSET ?
        """,
        [*params, per_page, offset],
    ).fetchall()

    total_collection = db.execute(
        "SELECT COALESCE(SUM(amount), 0) FROM payments WHERE UPPER(COALESCE(status, 'PENDING'))='DONE'"
    ).fetchone()[0]
    total_expense = db.execute("SELECT COALESCE(SUM(amount), 0) FROM expenses").fetchone()[0]
    balance = float(total_collection or 0) - float(total_expense or 0)

    return render_template(
        "expenses.html",
        active_page="expenses",
        rows=rows,
        total_collection=float(total_collection or 0),
        total_expense=float(total_expense or 0),
        balance=balance,
        search_item=search_item,
        search_date=search_date,
        sort=sort,
        page=page,
        total_pages=total_pages,
    )


@app.route("/expenses/new", methods=["GET", "POST"])
def expenses_new():
    auth_redirect = require_login()
    if auth_redirect:
        return auth_redirect
    db = get_db()
    if request.method == "POST":
        transaction_date_raw = request.form.get("transaction_date", "").strip()
        item_name = request.form.get("item_name", "").strip()
        amount = float(request.form.get("amount", "0") or 0)
        payment_mode = request.form.get("payment_mode", "").strip().upper() or None
        paid_by = request.form.get("paid_by", "").strip() or None

        if not transaction_date_raw or not item_name or amount <= 0:
            flash("Please fill transaction date, item name and valid amount.", "error")
            return redirect(url_for("expenses_new"))

        transaction_date = _normalize_date_for_storage(transaction_date_raw)
        if not transaction_date:
            flash("Invalid transaction date.", "error")
            return redirect(url_for("expenses_new"))

        bill_image_path = None
        bill_file = request.files.get("bill")
        if bill_file and bill_file.filename:
            safe = secure_filename(bill_file.filename)
            ext = Path(safe).suffix.lower()
            if ext not in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".pdf"}:
                flash("Bill must be an image or PDF file.", "error")
                return redirect(url_for("expenses_new"))
            unique_name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{safe}"
            out_path = UPLOAD_DIR / unique_name
            bill_file.save(out_path)
            bill_image_path = f"uploads/bills/{unique_name}"

        db.execute(
            """
            INSERT INTO expenses(transaction_date, item_name, amount, payment_mode, paid_by, bill_image_path, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            """,
            (transaction_date, item_name, amount, payment_mode, paid_by, bill_image_path),
        )
        db.commit()
        flash("Expense added successfully.", "success")
        return redirect(url_for("expenses"))

    return render_template(
        "expenses_new.html",
        active_page="expenses",
        today_iso=datetime.now().strftime("%Y-%m-%d"),
    )


@app.route("/expenses/<int:expense_id>/edit", methods=["GET", "POST"])
def expenses_edit(expense_id: int):
    auth_redirect = require_login()
    if auth_redirect:
        return auth_redirect
    db = get_db()
    row = db.execute(
        "SELECT id, transaction_date, item_name, amount, payment_mode, paid_by, bill_image_path FROM expenses WHERE id=?",
        (expense_id,),
    ).fetchone()
    if not row:
        flash("Expense record not found.", "error")
        return redirect(url_for("expenses"))

    if request.method == "POST":
        action = (request.form.get("action") or "save").strip().lower()
        if action == "delete":
            db.execute("DELETE FROM expenses WHERE id=?", (expense_id,))
            db.commit()
            flash("Expense deleted successfully.", "success")
            return redirect(url_for("expenses"))

        transaction_date_raw = request.form.get("transaction_date", "").strip()
        item_name = request.form.get("item_name", "").strip()
        amount = float(request.form.get("amount", "0") or 0)
        payment_mode = request.form.get("payment_mode", "").strip().upper() or None
        paid_by = request.form.get("paid_by", "").strip() or None
        existing_bill = request.form.get("existing_bill_path", "").strip() or None

        if not transaction_date_raw or not item_name or amount <= 0:
            flash("Please fill transaction date, item name and valid amount.", "error")
            return redirect(url_for("expenses_edit", expense_id=expense_id))

        transaction_date = _normalize_date_for_storage(transaction_date_raw)
        if not transaction_date:
            flash("Invalid transaction date.", "error")
            return redirect(url_for("expenses_edit", expense_id=expense_id))

        bill_image_path = existing_bill
        bill_file = request.files.get("bill")
        if bill_file and bill_file.filename:
            safe = secure_filename(bill_file.filename)
            ext = Path(safe).suffix.lower()
            if ext not in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".pdf"}:
                flash("Bill must be an image or PDF file.", "error")
                return redirect(url_for("expenses_edit", expense_id=expense_id))
            unique_name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{safe}"
            out_path = UPLOAD_DIR / unique_name
            bill_file.save(out_path)
            bill_image_path = f"uploads/bills/{unique_name}"

        db.execute(
            """
            UPDATE expenses
            SET transaction_date=?, item_name=?, amount=?, payment_mode=?, paid_by=?, bill_image_path=?, updated_at=CURRENT_TIMESTAMP
            WHERE id=?
            """,
            (transaction_date, item_name, amount, payment_mode, paid_by, bill_image_path, expense_id),
        )
        db.commit()
        flash("Expense updated successfully.", "success")
        return redirect(url_for("expenses"))

    return render_template(
        "expenses_edit.html",
        active_page="expenses",
        row=row,
        transaction_date_iso=_format_db_date_to_iso(row["transaction_date"]),
    )


@app.route("/expenses/export/excel")
def expenses_export_excel():
    db = get_db()
    search_item = request.args.get("item", "").strip()
    search_date = request.args.get("transaction_date", "").strip()
    sort = (request.args.get("sort", "desc") or "desc").lower()
    order = "DESC" if sort != "asc" else "ASC"

    where = " WHERE 1=1 "
    params = []
    if search_item:
        where += " AND item_name LIKE ? "
        params.append(f"%{search_item}%")
    if search_date:
        db_date = _normalize_date_for_storage(search_date)
        if db_date:
            where += " AND transaction_date = ? "
            params.append(db_date)

    rows = db.execute(
        f"""
        SELECT transaction_date, item_name, amount, payment_mode, paid_by, bill_image_path
        FROM expenses
        {where}
        ORDER BY transaction_date {order}, id {order}
        """,
        params,
    ).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(["Transaction Date", "Item Name", "Amount", "Payment Mode", "Paid By", "Bill"])
    for r in rows:
        ws.append([
            r["transaction_date"],
            r["item_name"],
            float(r["amount"] or 0),
            r["payment_mode"] or "",
            r["paid_by"] or "",
            r["bill_image_path"] or "",
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"expenses_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/health")
def api_health():
    return jsonify({"ok": True, "service": "society-api"})


@app.route("/api/login", methods=["POST"])
def api_login():
    db = get_db()
    payload = request.get_json(silent=True) or {}
    username = str(payload.get("username") or payload.get("identifier") or "").strip()
    password = str(payload.get("password") or "")
    if not username or not password:
        return jsonify({"ok": False, "error": "Email/mobile and password are required"}), 400

    if username == LOGIN_USERNAME and password == LOGIN_PASSWORD:
        admin_row = db.execute("SELECT id, name, mobile, email, role, block, flat FROM users WHERE LOWER(email)=LOWER(?)", ("admin@societies.app",)).fetchone()
        admin_id = int(admin_row["id"]) if admin_row else 1
        token = _create_api_token(admin_id, "ADMIN", LOGIN_USERNAME)
        return jsonify({"ok": True, "token": token, "user": {
            "id": admin_id, "name": "Admin", "mobile": "", "email": "admin@societies.app", "role": "ADMIN", "block": "N/A", "flat": "N/A"
        }})

    row = db.execute(
        """
        SELECT id, name, mobile, email, role, block, flat, status, password_hash
        FROM users
        WHERE (LOWER(email)=LOWER(?) OR mobile=?)
        LIMIT 1
        """,
        (username, username),
    ).fetchone()
    if not row or str(row["status"] or "").upper() != "APPROVED" or not check_password_hash(row["password_hash"], password):
        return jsonify({"ok": False, "error": "Invalid credentials or account not approved"}), 401

    token = _create_api_token(int(row["id"]), str(row["role"] or "TENANT").upper(), username)
    return jsonify({"ok": True, "token": token, "user": {
        "id": int(row["id"]),
        "name": row["name"],
        "mobile": row["mobile"],
        "email": row["email"],
        "role": str(row["role"] or "TENANT").upper(),
        "block": row["block"],
        "flat": row["flat"],
    }})


@app.route("/api/me")
def api_me():
    auth = require_api_auth()
    if auth:
        return auth
    user = _get_current_user_row()
    if not user:
        return jsonify({"ok": False, "error": "Unauthorized"}), 401
    return jsonify({"ok": True, "user": {
        "id": int(user["id"]),
        "name": user["name"] or "",
        "mobile": user["mobile"] or "",
        "email": user["email"] or "",
        "role": str(user["role"] or "TENANT").upper(),
        "block": user["block"] or "",
        "flat": user["flat"] or "",
        "vehicle_list": user.get("vehicle_list", "") or "",
        "photo_url": user.get("photo_url", "") or "",
        "living_from": _format_db_date_to_iso(user.get("living_from")),
        "rent_document_path": user.get("rent_document_path", "") or "",
        "id_card_document_path": user.get("id_card_document_path", "") or "",
    }})


@app.route("/api/me", methods=["PUT"])
def api_update_me():
    auth = require_api_auth()
    if auth:
        return auth
    user = _get_current_user_row()
    if not user:
        return jsonify({"ok": False, "error": "Unauthorized"}), 401
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    mobile = str(payload.get("mobile") or "").strip()
    email = str(payload.get("email") or "").strip().lower()
    vehicle_list = str(payload.get("vehicle_list") or "").strip() or None
    photo_url = str(payload.get("photo_url") or "").strip() or None
    living_from = _normalize_date_for_storage(payload.get("living_from"))
    if not name or not mobile or not email:
        return jsonify({"ok": False, "error": "name, mobile and email are required"}), 400
    db = get_db()
    db.execute(
        "UPDATE users SET name=?, mobile=?, email=?, vehicle_list=?, photo_url=?, living_from=?, updated_at=CURRENT_TIMESTAMP WHERE id=?",
        (name, mobile, email, vehicle_list, photo_url, living_from, int(user["id"])),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/change-password", methods=["POST"])
def api_change_password():
    auth = require_api_auth()
    if auth:
        return auth
    user = _get_current_user_row()
    payload = request.get_json(silent=True) or {}
    current_password = str(payload.get("current_password") or "")
    new_password = str(payload.get("new_password") or "")
    if not current_password or not new_password:
        return jsonify({"ok": False, "error": "current_password and new_password are required"}), 400
    if str(user["role"] or "").upper() == "ADMIN" and current_password == LOGIN_PASSWORD:
        return jsonify({"ok": False, "error": "Admin password is managed by server environment"}), 400
    if not check_password_hash(user["password_hash"], current_password):
        return jsonify({"ok": False, "error": "Current password is incorrect"}), 400
    db = get_db()
    db.execute("UPDATE users SET password_hash=?, updated_at=CURRENT_TIMESTAMP WHERE id=?", (generate_password_hash(new_password), int(user["id"])))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/forgot-password", methods=["POST"])
def api_forgot_password():
    payload = request.get_json(silent=True) or {}
    identifier = str(payload.get("identifier") or "").strip()
    if not identifier:
        return jsonify({"ok": False, "error": "identifier is required"}), 400
    return jsonify({"ok": True, "message": "OTP flow pending implementation"})


@app.route("/api/register-requests", methods=["POST"])
def api_create_register_request():
    db = get_db()
    if request.content_type and request.content_type.startswith("multipart/form-data"):
        payload = request.form
        rent_doc = request.files.get("rent_document")
        id_doc = request.files.get("id_card_document")
    else:
        payload = request.get_json(silent=True) or {}
        rent_doc = None
        id_doc = None
    name = str(payload.get("name") or "").strip()
    mobile = str(payload.get("mobile") or "").strip()
    email = str(payload.get("email") or "").strip().lower()
    password = str(payload.get("password") or "")
    block = str(payload.get("block") or "").strip()
    flat = str(payload.get("flat") or "").strip()
    living_from = _normalize_date_for_storage(payload.get("living_from"))
    if not all([name, mobile, email, password, block, flat]):
        return jsonify({"ok": False, "error": "name, mobile, email, password, block, flat are required"}), 400
    if block not in ALLOWED_BLOCKS or flat not in ALLOWED_FLATS:
        return jsonify({"ok": False, "error": "Invalid block/flat"}), 400

    exists_user = db.execute("SELECT id FROM users WHERE LOWER(email)=LOWER(?) OR mobile=?", (email, mobile)).fetchone()
    exists_request = db.execute("SELECT id FROM registration_requests WHERE LOWER(email)=LOWER(?) OR mobile=?", (email, mobile)).fetchone()
    if exists_user or exists_request:
        return jsonify({"ok": False, "error": "A user/request already exists with this email or mobile"}), 400

    rent_document_path, rent_err = _save_uploaded_kyc(rent_doc)
    if rent_err:
        return jsonify({"ok": False, "error": rent_err}), 400
    id_card_document_path, id_err = _save_uploaded_kyc(id_doc)
    if id_err:
        return jsonify({"ok": False, "error": id_err}), 400

    try:
        db.execute(
            """
            INSERT INTO registration_requests(
              name, mobile, email, password_hash, block, flat, living_from, rent_document_path, id_card_document_path, status, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'PENDING', CURRENT_TIMESTAMP)
            """,
            (name, mobile, email, generate_password_hash(password), block, flat, living_from, rent_document_path, id_card_document_path),
        )
        db.commit()
        return jsonify({"ok": True})
    except Exception:
        return jsonify({"ok": False, "error": "Unable to create registration request"}), 500


@app.route("/api/register-requests")
def api_list_register_requests():
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    rows = db.execute(
        """
        SELECT id, name, mobile, email, block, flat, living_from, rent_document_path, id_card_document_path, status, created_at
        FROM registration_requests
        ORDER BY created_at DESC, id DESC
        """
    ).fetchall()
    return jsonify({"ok": True, "data": [{
        "id": int(r["id"]),
        "name": r["name"],
        "mobile": r["mobile"],
        "email": r["email"],
        "block": r["block"],
        "flat": r["flat"],
        "living_from": _format_db_date_to_iso(r.get("living_from")),
        "rent_document_path": r.get("rent_document_path") or "",
        "id_card_document_path": r.get("id_card_document_path") or "",
        "status": r["status"],
        "created_at": r["created_at"],
    } for r in rows]})


@app.route("/api/register-requests/<int:request_id>/approve", methods=["POST"])
def api_approve_register_request(request_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    payload = request.get_json(silent=True) or {}
    role = str(payload.get("role") or "").strip().upper()
    row = db.execute("SELECT * FROM registration_requests WHERE id=?", (request_id,)).fetchone()
    if not row:
        return jsonify({"ok": False, "error": "Request not found"}), 404
    block = str(payload.get("block") or row["block"]).strip()
    flat = str(payload.get("flat") or row["flat"]).strip()
    if role not in {"ADMIN", "OWNER", "TENANT"}:
        return jsonify({"ok": False, "error": "Invalid role"}), 400
    if block and block not in ALLOWED_BLOCKS:
        return jsonify({"ok": False, "error": "Invalid block"}), 400
    if flat and flat not in ALLOWED_FLATS:
        return jsonify({"ok": False, "error": "Invalid flat"}), 400
    if str(row["status"] or "").upper() != "PENDING":
        return jsonify({"ok": False, "error": "Only pending requests can be approved"}), 400
    try:
        # Ensure flat exists and map approved user to that flat's owner/tenant details.
        db.execute(
            """
            INSERT INTO properties(block, flat, updated_at)
            VALUES (?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(block, flat) DO UPDATE SET updated_at=CURRENT_TIMESTAMP
            """,
            (block, flat),
        )
        property_row = db.execute("SELECT id FROM properties WHERE block=? AND flat=?", (block, flat)).fetchone()
        property_id = int(property_row[0]) if property_row else None

        db.execute(
            """
            INSERT INTO users(
              name, mobile, email, password_hash, role, block, flat, status,
              living_from, rent_document_path, id_card_document_path, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, 'APPROVED', ?, ?, ?, CURRENT_TIMESTAMP)
            """,
            (
                row["name"], row["mobile"], row["email"], row["password_hash"], role, block, flat,
                row.get("living_from"), row.get("rent_document_path"), row.get("id_card_document_path"),
            ),
        )

        if property_id and role == "OWNER":
            db.execute(
                """
                INSERT INTO owner_details(property_id, owner_name, owner_contact, is_occupied, occupied_by, updated_at)
                VALUES (?, ?, ?, 0, 'OWNER', CURRENT_TIMESTAMP)
                ON CONFLICT(property_id) DO UPDATE SET
                  owner_name=excluded.owner_name,
                  owner_contact=excluded.owner_contact,
                  updated_at=CURRENT_TIMESTAMP
                """,
                (property_id, row["name"], row["mobile"]),
            )
        elif property_id and role == "TENANT":
            db.execute(
                """
                INSERT INTO owner_details(
                  property_id, is_occupied, occupied_by, tenant_name, tenant_contact, tenant_living_from, updated_at
                )
                VALUES (?, 1, 'TENANT', ?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(property_id) DO UPDATE SET
                  is_occupied=1,
                  occupied_by='TENANT',
                  tenant_name=excluded.tenant_name,
                  tenant_contact=excluded.tenant_contact,
                  tenant_living_from=excluded.tenant_living_from,
                  updated_at=CURRENT_TIMESTAMP
                """,
                (property_id, row["name"], row["mobile"], row.get("living_from")),
            )

        db.execute("UPDATE registration_requests SET status='APPROVED', block=?, flat=?, updated_at=CURRENT_TIMESTAMP WHERE id=?", (block, flat, request_id))
        db.commit()
        return jsonify({"ok": True})
    except Exception:
        return jsonify({"ok": False, "error": "Unable to approve request"}), 500


@app.route("/api/register-requests/<int:request_id>/reject", methods=["POST"])
def api_reject_register_request(request_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    row = db.execute("SELECT id FROM registration_requests WHERE id=?", (request_id,)).fetchone()
    if not row:
        return jsonify({"ok": False, "error": "Request not found"}), 404
    db.execute("UPDATE registration_requests SET status='REJECTED', updated_at=CURRENT_TIMESTAMP WHERE id=?", (request_id,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/users")
def api_users():
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    rows = db.execute(
        "SELECT id, name, mobile, email, role, block, flat, status, created_at FROM users ORDER BY id DESC"
    ).fetchall()
    return jsonify({"ok": True, "data": [{
        "id": int(r["id"]),
        "name": r["name"],
        "mobile": r["mobile"],
        "email": r["email"],
        "role": str(r["role"] or "TENANT").upper(),
        "block": r["block"],
        "flat": r["flat"],
        "status": r["status"],
        "created_at": r["created_at"],
    } for r in rows]})


@app.route("/api/users/<int:user_id>", methods=["PUT"])
def api_update_user(user_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    payload = request.get_json(silent=True) or {}
    role = str(payload.get("role") or "").strip().upper()
    block = str(payload.get("block") or "").strip()
    flat = str(payload.get("flat") or "").strip()
    if role not in {"ADMIN", "OWNER", "TENANT"}:
        return jsonify({"ok": False, "error": "Invalid role"}), 400
    if not block or not flat:
        return jsonify({"ok": False, "error": "block and flat are required"}), 400
    db.execute("UPDATE users SET role=?, block=?, flat=?, updated_at=CURRENT_TIMESTAMP WHERE id=?", (role, block, flat, user_id))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/dashboard")
def api_dashboard():
    db = get_db()
    summary = _build_public_summary(db)
    return jsonify({"ok": True, "data": summary})


@app.route("/api/payments")
def api_payments():
    db = get_db()
    year = int(request.args.get("year", datetime.now().year))
    month = int(request.args.get("month", datetime.now().month))
    block = request.args.get("block", "ALL")
    flat = request.args.get("flat", "").strip()
    payment_date = _normalize_date_for_storage(request.args.get("payment_date"))
    scope = (request.args.get("scope", "month") or "month").strip().lower()

    query = """
    SELECT p.id AS property_id, p.block, p.flat, pay.id AS payment_id, pay.amount, pay.status, pay.year, pay.month, pay.payment_date, pay.mode_of_payment, pay.notes,
           od.tenant_name
    FROM payments pay
    JOIN properties p ON p.id = pay.property_id
    LEFT JOIN owner_details od ON od.property_id = p.id
    WHERE UPPER(COALESCE(pay.status, 'PENDING')) = 'DONE'
    """
    params = []
    if block != "ALL":
        query += " AND p.block = ?"
        params.append(block)
    if flat:
        query += " AND p.flat LIKE ?"
        params.append(f"%{flat}%")
    if payment_date:
        query += " AND pay.payment_date = ?"
        params.append(payment_date)
    elif scope == "today":
        today_db = datetime.now().strftime("%m/%d/%Y")
        query += " AND pay.payment_date = ?"
        params.append(today_db)
    elif scope == "month":
        query += " AND pay.year=? AND pay.month=?"
        params.extend([year, month])
    elif scope == "year":
        query += " AND pay.year=?"
        params.append(year)
    query += " ORDER BY pay.updated_at DESC, pay.id DESC"
    rows = db.execute(query, params).fetchall()

    entries = []
    total_amount = 0.0
    for r in rows:
        amount = float(r["amount"] or 0)
        total_amount += amount
        entries.append({
            "payment_id": r["payment_id"],
            "property_id": r["property_id"],
            "block": r["block"],
            "flat": r["flat"],
            "amount": amount,
            "status": (r["status"] or "DONE").upper(),
            "year": int(r["year"] or 0),
            "month": int(r["month"] or 0),
            "payment_date": _format_db_date_to_iso(r["payment_date"]),
            "mode_of_payment": (r["mode_of_payment"] or "").upper(),
            "notes": r["notes"] or "",
            "tenant_name": r["tenant_name"] or "",
        })

    totals_row = db.execute(
        "SELECT COALESCE(SUM(amount), 0) FROM payments WHERE UPPER(COALESCE(status, 'PENDING'))='DONE'"
    ).fetchone()
    grand_total = float(totals_row[0] or 0)

    return jsonify({
        "ok": True,
        "data": {
            "year": year,
            "month": month,
            "scope": scope,
            "entries": entries,
            "count": len(entries),
            "total_amount": total_amount,
            "grand_total": grand_total,
            "months": MONTHS,
        },
    })


@app.route("/api/payments", methods=["POST"])
def api_create_or_update_payment():
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    payload = request.get_json(silent=True) or {}

    block = str(payload.get("block") or "").strip()
    flat = str(payload.get("flat") or "").strip()
    year = int(payload.get("year") or 0)
    month = int(payload.get("month") or 0)
    amount = float(payload.get("amount") or 0)
    payment_date = _normalize_date_for_storage(payload.get("payment_date"))
    notes = str(payload.get("notes") or "").strip() or None
    mode_of_payment = str(payload.get("mode_of_payment") or "").strip().upper() or None
    received_by = str(payload.get("received_by") or "").strip() or None

    if block not in ALLOWED_BLOCKS or flat not in ALLOWED_FLATS or year < 2000 or month < 1 or month > 12:
        return jsonify({"ok": False, "error": "Invalid block/flat/year/month"}), 400
    if not payment_date:
        return jsonify({"ok": False, "error": "Invalid payment_date. Use YYYY-MM-DD"}), 400
    if mode_of_payment not in {"CASH", "ONLINE", None}:
        return jsonify({"ok": False, "error": "mode_of_payment must be CASH or ONLINE"}), 400
    if mode_of_payment == "CASH" and not received_by:
        return jsonify({"ok": False, "error": "received_by is required for CASH payments"}), 400

    db.execute(
        """
        INSERT INTO properties(block, flat, updated_at)
        VALUES (?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(block, flat) DO UPDATE SET updated_at=CURRENT_TIMESTAMP
        """,
        (block, flat),
    )
    property_id = db.execute(
        "SELECT id FROM properties WHERE block=? AND flat=?",
        (block, flat),
    ).fetchone()[0]

    status = "DONE"
    db.execute(
        """
        INSERT INTO payments(property_id, year, month, amount, payment_date, status, notes, source, mode_of_payment, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, 'mobile-api', ?, CURRENT_TIMESTAMP)
        ON CONFLICT(property_id, year, month) DO UPDATE SET
          amount=excluded.amount,
          payment_date=excluded.payment_date,
          status=excluded.status,
          notes=excluded.notes,
          mode_of_payment=excluded.mode_of_payment,
          source='mobile-api',
          updated_at=CURRENT_TIMESTAMP
        """,
        (property_id, year, month, amount, payment_date, status, notes, mode_of_payment),
    )
    if received_by:
        db.execute(
            "UPDATE payments SET notes=TRIM(COALESCE(notes,'') || CASE WHEN COALESCE(notes,'')='' THEN '' ELSE ' | ' END || ?) WHERE property_id=? AND year=? AND month=?",
            (f"Received By: {received_by}", property_id, year, month),
        )
    db.commit()
    return jsonify({"ok": True, "property_id": property_id})


@app.route("/api/payments/<int:property_id>/<int:year>/<int:month>", methods=["PUT"])
def api_update_payment_by_key(property_id: int, year: int, month: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    payload = request.get_json(silent=True) or {}
    amount = float(payload.get("amount") or 0)
    payment_date = _normalize_date_for_storage(payload.get("payment_date"))
    notes = str(payload.get("notes") or "").strip() or None
    mode_of_payment = str(payload.get("mode_of_payment") or "").strip().upper() or None
    received_by = str(payload.get("received_by") or "").strip() or None
    status = str(payload.get("status") or "DONE").strip().upper()
    if status not in {"DONE", "PENDING", "LOCKED"}:
        return jsonify({"ok": False, "error": "Invalid status"}), 400
    if mode_of_payment not in {"CASH", "ONLINE", None}:
        return jsonify({"ok": False, "error": "mode_of_payment must be CASH or ONLINE"}), 400
    if mode_of_payment == "CASH" and not received_by:
        return jsonify({"ok": False, "error": "received_by is required for CASH payments"}), 400
    existing = db.execute(
        "SELECT id FROM payments WHERE property_id=? AND year=? AND month=?",
        (property_id, year, month),
    ).fetchone()
    if not existing:
        return jsonify({"ok": False, "error": "Payment not found"}), 404
    db.execute(
        """
        UPDATE payments
        SET amount=?, payment_date=?, status=?, notes=?, mode_of_payment=?, source='mobile-api', updated_at=CURRENT_TIMESTAMP
        WHERE property_id=? AND year=? AND month=?
        """,
        (amount if status == "DONE" else 0, payment_date, status, notes, mode_of_payment, property_id, year, month),
    )
    if received_by:
        db.execute(
            "UPDATE payments SET notes=TRIM(COALESCE(notes,'') || CASE WHEN COALESCE(notes,'')='' THEN '' ELSE ' | ' END || ?) WHERE property_id=? AND year=? AND month=?",
            (f"Received By: {received_by}", property_id, year, month),
        )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/payments/<int:property_id>/<int:year>/<int:month>", methods=["DELETE"])
def api_delete_payment_by_key(property_id: int, year: int, month: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    db.execute(
        "DELETE FROM payments WHERE property_id=? AND year=? AND month=?",
        (property_id, year, month),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/expenses")
def api_expenses():
    db = get_db()
    search_item = request.args.get("item", "").strip()
    query_mode = (request.args.get("scope", "all") or "all").strip().lower()
    query = """
    SELECT id, transaction_date, item_name, quantity, amount, payment_mode, paid_by, bill_image_path, created_at
    FROM expenses
    WHERE 1=1
    """
    params = []
    if search_item:
        query += " AND item_name LIKE ? "
        params.append(f"%{search_item}%")
    if query_mode == "month":
        now = datetime.now()
        query += " AND strftime('%Y', substr(transaction_date, 7, 4) || '-' || substr(transaction_date, 1, 2) || '-' || substr(transaction_date, 4, 2)) = ? "
        query += " AND strftime('%m', substr(transaction_date, 7, 4) || '-' || substr(transaction_date, 1, 2) || '-' || substr(transaction_date, 4, 2)) = ? "
        params.extend([f"{now.year:04d}", f"{now.month:02d}"])
    query += " ORDER BY transaction_date DESC, id DESC LIMIT 100"
    rows = db.execute(query, params).fetchall()
    data = []
    for row in rows:
        data.append({
            "id": row["id"],
            "transaction_date": _format_db_date_to_iso(row["transaction_date"]),
            "item_name": row["item_name"],
            "quantity": row["quantity"] or "",
            "amount": float(row["amount"] or 0),
            "payment_mode": row["payment_mode"] or "",
            "paid_by": row["paid_by"] or "",
            "bill_image_path": row["bill_image_path"] or "",
        })
    total_collection = float(db.execute(
        "SELECT COALESCE(SUM(amount), 0) FROM payments WHERE UPPER(COALESCE(status, 'PENDING'))='DONE'"
    ).fetchone()[0] or 0)
    total_expense = float(db.execute("SELECT COALESCE(SUM(amount), 0) FROM expenses").fetchone()[0] or 0)
    return jsonify({"ok": True, "data": data, "summary": {"total_collection": total_collection, "total_expense": total_expense, "balance": total_collection - total_expense}})


@app.route("/api/expenses", methods=["POST"])
def api_create_expense():
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    if request.content_type and request.content_type.startswith("multipart/form-data"):
        raw = request.form
        bill_file = request.files.get("bill")
    else:
        raw = request.get_json(silent=True) or {}
        bill_file = None
    transaction_date = _normalize_date_for_storage(raw.get("transaction_date"))
    item_name = str(raw.get("item_name") or "").strip()
    quantity = str(raw.get("quantity") or "").strip() or None
    amount = float(raw.get("amount") or 0)
    payment_mode = str(raw.get("payment_mode") or "").strip().upper() or None
    paid_by = str(raw.get("paid_by") or "").strip() or None
    if not transaction_date or not item_name or amount <= 0:
        return jsonify({"ok": False, "error": "Invalid transaction_date/item_name/amount"}), 400
    bill_image_path, bill_error = _save_uploaded_bill(bill_file)
    if bill_error:
        return jsonify({"ok": False, "error": bill_error}), 400
    db.execute(
        """
        INSERT INTO expenses(transaction_date, item_name, quantity, amount, payment_mode, paid_by, bill_image_path, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        """,
        (transaction_date, item_name, quantity, amount, payment_mode, paid_by, bill_image_path),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/expenses/<int:expense_id>", methods=["PUT"])
def api_update_expense(expense_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    existing = db.execute("SELECT id, bill_image_path FROM expenses WHERE id=?", (expense_id,)).fetchone()
    if not existing:
        return jsonify({"ok": False, "error": "Expense not found"}), 404
    if request.content_type and request.content_type.startswith("multipart/form-data"):
        raw = request.form
        bill_file = request.files.get("bill")
    else:
        raw = request.get_json(silent=True) or {}
        bill_file = None
    transaction_date = _normalize_date_for_storage(raw.get("transaction_date"))
    item_name = str(raw.get("item_name") or "").strip()
    quantity = str(raw.get("quantity") or "").strip() or None
    amount = float(raw.get("amount") or 0)
    payment_mode = str(raw.get("payment_mode") or "").strip().upper() or None
    paid_by = str(raw.get("paid_by") or "").strip() or None
    if not transaction_date or not item_name or amount <= 0:
        return jsonify({"ok": False, "error": "Invalid transaction_date/item_name/amount"}), 400
    bill_image_path = str(raw.get("existing_bill_path") or existing["bill_image_path"] or "").strip() or None
    uploaded_bill, bill_error = _save_uploaded_bill(bill_file)
    if bill_error:
        return jsonify({"ok": False, "error": bill_error}), 400
    if uploaded_bill:
        bill_image_path = uploaded_bill
    db.execute(
        """
        UPDATE expenses
        SET transaction_date=?, item_name=?, quantity=?, amount=?, payment_mode=?, paid_by=?, bill_image_path=?, updated_at=CURRENT_TIMESTAMP
        WHERE id=?
        """,
        (transaction_date, item_name, quantity, amount, payment_mode, paid_by, bill_image_path, expense_id),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/expenses/<int:expense_id>", methods=["DELETE"])
def api_delete_expense(expense_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    db.execute("DELETE FROM expenses WHERE id=?", (expense_id,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/owners")
def api_owners():
    db = get_db()
    block = (request.args.get("block") or "ALL").strip()
    flat = (request.args.get("flat") or "").strip()
    owner_query = (request.args.get("owner") or "").strip()
    contact_query = (request.args.get("contact") or "").strip()

    query = """
    SELECT p.id AS property_id, p.block, p.flat, od.owner_name, od.owner_contact, od.is_occupied, od.occupied_by, od.tenant_name
    FROM properties p
    LEFT JOIN owner_details od ON od.property_id = p.id
    WHERE 1=1
    """
    params = []
    if block != "ALL":
        query += " AND p.block = ?"
        params.append(block)
    if flat:
        query += " AND p.flat LIKE ?"
        params.append(f"%{flat}%")
    if owner_query:
        query += " AND COALESCE(od.owner_name, '') LIKE ?"
        params.append(f"%{owner_query}%")
    if contact_query:
        query += " AND COALESCE(od.owner_contact, '') LIKE ?"
        params.append(f"%{contact_query}%")
    query += " ORDER BY p.block, p.flat"
    rows = db.execute(query, params).fetchall()
    data = [{
        "property_id": r["property_id"],
        "block": r["block"],
        "flat": r["flat"],
        "owner_name": r["owner_name"] or "",
        "owner_contact": r["owner_contact"] or "",
        "is_occupied": int(r["is_occupied"] or 0),
        "occupied_by": (r["occupied_by"] or "OWNER").upper(),
        "tenant_name": r["tenant_name"] or "",
    } for r in rows]
    return jsonify({"ok": True, "data": data})


@app.route("/api/owners/<int:property_id>")
def api_owner_detail(property_id: int):
    db = get_db()
    row = db.execute(
        """
        SELECT p.id AS property_id, p.block, p.flat,
               od.owner_name, od.owner_contact, od.is_occupied, od.occupied_by,
               od.tenant_name, od.tenant_contact, od.tenant_vehicle_list,
               od.tenant_photo_url, od.tenant_living_from, od.tenant_guard_payment_details
        FROM properties p
        LEFT JOIN owner_details od ON od.property_id = p.id
        WHERE p.id=?
        """,
        (property_id,),
    ).fetchone()
    if not row:
        return jsonify({"ok": False, "error": "Flat not found"}), 404
    tenant_name = row["tenant_name"] or ""
    tenant_contact = row["tenant_contact"] or ""
    tenant_vehicle_list = row["tenant_vehicle_list"] or ""
    tenant_photo_url = row["tenant_photo_url"] or ""
    tenant_living_from = _format_db_date_to_iso(row["tenant_living_from"])
    tenant_guard_payment_details = row["tenant_guard_payment_details"] or ""
    if not tenant_name:
        fallback_tenant = db.execute(
            """
            SELECT od.tenant_name, od.tenant_contact, od.tenant_vehicle_list, od.tenant_photo_url, od.tenant_living_from, od.tenant_guard_payment_details
            FROM owner_details od
            JOIN properties p2 ON p2.id = od.property_id
            WHERE p2.block=? AND p2.flat=? AND UPPER(COALESCE(od.occupied_by, 'OWNER'))='TENANT' AND COALESCE(od.tenant_name, '') <> ''
            ORDER BY od.updated_at DESC, od.id DESC
            LIMIT 1
            """,
            (row["block"], row["flat"]),
        ).fetchone()
        if fallback_tenant:
            tenant_name = fallback_tenant["tenant_name"] or ""
            tenant_contact = fallback_tenant["tenant_contact"] or ""
            tenant_vehicle_list = fallback_tenant["tenant_vehicle_list"] or ""
            tenant_photo_url = fallback_tenant["tenant_photo_url"] or ""
            tenant_living_from = _format_db_date_to_iso(fallback_tenant["tenant_living_from"])
            tenant_guard_payment_details = fallback_tenant["tenant_guard_payment_details"] or ""

    return jsonify({"ok": True, "data": {
        "property_id": row["property_id"],
        "block": row["block"],
        "flat": row["flat"],
        "owner_name": row["owner_name"] or "",
        "owner_contact": row["owner_contact"] or "",
        "is_occupied": int(row["is_occupied"] or 0),
        "occupied_by": (row["occupied_by"] or "OWNER").upper(),
        "tenant_name": tenant_name,
        "tenant_contact": tenant_contact,
        "tenant_vehicle_list": tenant_vehicle_list,
        "tenant_photo_url": tenant_photo_url,
        "tenant_living_from": tenant_living_from,
        "tenant_guard_payment_details": tenant_guard_payment_details,
    }})


@app.route("/api/owners/<int:property_id>", methods=["PUT"])
def api_owner_upsert(property_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    payload = request.get_json(silent=True) or {}
    exists = db.execute("SELECT id FROM properties WHERE id=?", (property_id,)).fetchone()
    if not exists:
        return jsonify({"ok": False, "error": "Flat not found"}), 404
    owner_name = str(payload.get("owner_name") or "").strip() or None
    owner_contact = str(payload.get("owner_contact") or "").strip() or None
    is_occupied = 1 if int(payload.get("is_occupied") or 0) else 0
    occupied_by = str(payload.get("occupied_by") or "OWNER").strip().upper()
    if occupied_by not in {"OWNER", "TENANT"}:
        return jsonify({"ok": False, "error": "occupied_by must be OWNER or TENANT"}), 400
    current = db.execute("SELECT * FROM owner_details WHERE property_id=?", (property_id,)).fetchone()
    property_row = db.execute("SELECT block, flat FROM properties WHERE id=?", (property_id,)).fetchone()
    fallback_tenant = None
    if property_row:
        fallback_tenant = db.execute(
            """
            SELECT od.tenant_name, od.tenant_contact, od.tenant_vehicle_list, od.tenant_photo_url, od.tenant_living_from, od.tenant_guard_payment_details
            FROM owner_details od
            JOIN properties p2 ON p2.id = od.property_id
            WHERE p2.block=? AND p2.flat=? AND UPPER(COALESCE(od.occupied_by, 'OWNER'))='TENANT' AND COALESCE(od.tenant_name, '') <> ''
            ORDER BY od.updated_at DESC, od.id DESC
            LIMIT 1
            """,
            (property_row["block"], property_row["flat"]),
        ).fetchone()

    def _pick(key: str, normalize_date: bool = False):
        if key in payload:
            val = payload.get(key)
            if normalize_date:
                return _normalize_date_for_storage(val)
            return str(val or "").strip() or None
        if current and key in current:
            return current[key]
        if fallback_tenant and key in fallback_tenant:
            return fallback_tenant[key]
        return None

    tenant_name = _pick("tenant_name")
    tenant_contact = _pick("tenant_contact")
    tenant_vehicle_list = _pick("tenant_vehicle_list")
    tenant_photo_url = _pick("tenant_photo_url")
    tenant_living_from = _pick("tenant_living_from", normalize_date=True) if "tenant_living_from" in payload else (current["tenant_living_from"] if current and "tenant_living_from" in current else (fallback_tenant["tenant_living_from"] if fallback_tenant and "tenant_living_from" in fallback_tenant else None))
    tenant_guard_payment_details = _pick("tenant_guard_payment_details")
    db.execute(
        """
        INSERT INTO owner_details(
          property_id, owner_name, owner_contact, is_occupied, occupied_by, tenant_name, tenant_contact, tenant_vehicle_list, tenant_photo_url, tenant_living_from, tenant_guard_payment_details, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(property_id) DO UPDATE SET
          owner_name=excluded.owner_name,
          owner_contact=excluded.owner_contact,
          is_occupied=excluded.is_occupied,
          occupied_by=excluded.occupied_by,
          tenant_name=excluded.tenant_name,
          tenant_contact=excluded.tenant_contact,
          tenant_vehicle_list=excluded.tenant_vehicle_list,
          tenant_photo_url=excluded.tenant_photo_url,
          tenant_living_from=excluded.tenant_living_from,
          tenant_guard_payment_details=excluded.tenant_guard_payment_details,
          updated_at=CURRENT_TIMESTAMP
        """,
        (property_id, owner_name, owner_contact, is_occupied, occupied_by, tenant_name, tenant_contact, tenant_vehicle_list, tenant_photo_url, tenant_living_from, tenant_guard_payment_details),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/owners/<int:property_id>", methods=["DELETE"])
def api_owner_delete(property_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    db.execute("DELETE FROM owner_details WHERE property_id=?", (property_id,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/tenants")
def api_tenants():
    db = get_db()
    block = (request.args.get("block") or "ALL").strip()
    flat = (request.args.get("flat") or "").strip()
    tenant_query = (request.args.get("tenant") or "").strip()
    contact_query = (request.args.get("contact") or "").strip()
    query = """
    SELECT p.id AS property_id, p.block, p.flat,
           od.tenant_name, od.tenant_contact, od.tenant_vehicle_list,
           od.tenant_photo_url, od.tenant_living_from, od.tenant_guard_payment_details
    FROM properties p
    JOIN owner_details od ON od.property_id = p.id
    WHERE UPPER(COALESCE(od.occupied_by, 'OWNER'))='TENANT'
      AND COALESCE(od.tenant_name, '') <> ''
    """
    params = []
    if block != "ALL":
        query += " AND p.block = ?"
        params.append(block)
    if flat:
        query += " AND p.flat LIKE ?"
        params.append(f"%{flat}%")
    if tenant_query:
        query += " AND COALESCE(od.tenant_name, '') LIKE ?"
        params.append(f"%{tenant_query}%")
    if contact_query:
        query += " AND COALESCE(od.tenant_contact, '') LIKE ?"
        params.append(f"%{contact_query}%")
    query += " ORDER BY p.block, p.flat"
    rows = db.execute(query, params).fetchall()
    data = [{
        "property_id": r["property_id"],
        "block": r["block"],
        "flat": r["flat"],
        "tenant_name": r["tenant_name"] or "",
        "tenant_contact": r["tenant_contact"] or "",
        "tenant_vehicle_list": r["tenant_vehicle_list"] or "",
        "tenant_photo_url": r["tenant_photo_url"] or "",
        "tenant_living_from": _format_db_date_to_iso(r["tenant_living_from"]),
        "tenant_guard_payment_details": r["tenant_guard_payment_details"] or "",
    } for r in rows]
    return jsonify({"ok": True, "data": data})


@app.route("/api/tenants/<int:property_id>")
def api_tenant_detail(property_id: int):
    db = get_db()
    row = db.execute(
        """
        SELECT p.id AS property_id, p.block, p.flat,
               od.owner_name, od.owner_contact,
               od.tenant_name, od.tenant_contact, od.tenant_vehicle_list,
               od.tenant_photo_url, od.tenant_living_from, od.tenant_guard_payment_details
        FROM properties p
        JOIN owner_details od ON od.property_id = p.id
        WHERE p.id=?
        """,
        (property_id,),
    ).fetchone()
    if not row:
        return jsonify({"ok": False, "error": "Tenant record not found"}), 404
    payment_rows = db.execute(
        """
        SELECT pay.year, pay.month, pay.amount, pay.payment_date, pay.mode_of_payment
        FROM payments pay
        JOIN properties p2 ON p2.id = pay.property_id
        WHERE p2.block=? AND p2.flat=? AND UPPER(COALESCE(pay.status, 'PENDING'))='DONE'
        ORDER BY year DESC, month DESC
        """,
        (row["block"], row["flat"]),
    ).fetchall()
    payment_history = [{
        "year": int(pr["year"] or 0),
        "month": int(pr["month"] or 0),
        "amount": float(pr["amount"] or 0),
        "payment_date": _format_db_date_to_iso(pr["payment_date"]),
        "mode_of_payment": (pr["mode_of_payment"] or "").upper(),
    } for pr in payment_rows]

    owner_name = row["owner_name"] or ""
    owner_contact = row["owner_contact"] or ""
    if not owner_name:
        fallback_owner = db.execute(
            """
            SELECT od.owner_name, od.owner_contact
            FROM owner_details od
            JOIN properties p2 ON p2.id = od.property_id
            WHERE p2.block=? AND p2.flat=? AND COALESCE(od.owner_name, '') <> ''
            ORDER BY od.updated_at DESC, od.id DESC
            LIMIT 1
            """,
            (row["block"], row["flat"]),
        ).fetchone()
        if fallback_owner:
            owner_name = fallback_owner["owner_name"] or ""
            owner_contact = fallback_owner["owner_contact"] or ""

    return jsonify({"ok": True, "data": {
        "property_id": row["property_id"],
        "block": row["block"],
        "flat": row["flat"],
        "owner_name": owner_name,
        "owner_contact": owner_contact,
        "tenant_name": row["tenant_name"] or "",
        "tenant_contact": row["tenant_contact"] or "",
        "tenant_vehicle_list": row["tenant_vehicle_list"] or "",
        "tenant_photo_url": row["tenant_photo_url"] or "",
        "tenant_living_from": _format_db_date_to_iso(row["tenant_living_from"]),
        "tenant_guard_payment_details": row["tenant_guard_payment_details"] or "",
        "payment_history": payment_history,
    }})


@app.route("/api/owner-lookup")
def api_owner_lookup():
    db = get_db()
    block = (request.args.get("block") or "").strip()
    flat = (request.args.get("flat") or "").strip()
    if not block or not flat:
        return jsonify({"ok": True, "data": None})
    row = db.execute(
        """
        SELECT p.id AS property_id, p.block, p.flat, COALESCE(od.owner_name, '') AS owner_name
        FROM properties p
        LEFT JOIN owner_details od ON od.property_id = p.id
        WHERE p.block=? AND p.flat=?
        LIMIT 1
        """,
        (block, flat),
    ).fetchone()
    if not row:
        return jsonify({"ok": True, "data": None})
    return jsonify({"ok": True, "data": {
        "property_id": row["property_id"],
        "block": row["block"],
        "flat": row["flat"],
        "owner_name": row["owner_name"] or "",
    }})


@app.route("/api/tenants", methods=["POST"])
def api_create_tenant():
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    payload = request.get_json(silent=True) or {}
    block = str(payload.get("block") or "").strip()
    flat = str(payload.get("flat") or "").strip()
    if block not in ALLOWED_BLOCKS or flat not in ALLOWED_FLATS:
        return jsonify({"ok": False, "error": "Invalid block/flat"}), 400
    db.execute(
        """
        INSERT INTO properties(block, flat, updated_at)
        VALUES (?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(block, flat) DO UPDATE SET updated_at=CURRENT_TIMESTAMP
        """,
        (block, flat),
    )
    property_id = db.execute("SELECT id FROM properties WHERE block=? AND flat=?", (block, flat)).fetchone()[0]
    tenant_name = str(payload.get("tenant_name") or "").strip() or None
    tenant_contact = str(payload.get("tenant_contact") or "").strip() or None
    tenant_vehicle_list = str(payload.get("tenant_vehicle_list") or "").strip() or None
    tenant_photo_url = str(payload.get("tenant_photo_url") or "").strip() or None
    tenant_living_from = _normalize_date_for_storage(payload.get("tenant_living_from"))
    db.execute(
        """
        INSERT INTO owner_details(property_id, is_occupied, occupied_by, tenant_name, tenant_contact, tenant_vehicle_list, tenant_photo_url, tenant_living_from, updated_at)
        VALUES (?, 1, 'TENANT', ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(property_id) DO UPDATE SET
          is_occupied=1,
          occupied_by='TENANT',
          tenant_name=excluded.tenant_name,
          tenant_contact=excluded.tenant_contact,
          tenant_vehicle_list=excluded.tenant_vehicle_list,
          tenant_photo_url=excluded.tenant_photo_url,
          tenant_living_from=excluded.tenant_living_from,
          updated_at=CURRENT_TIMESTAMP
        """,
        (property_id, tenant_name, tenant_contact, tenant_vehicle_list, tenant_photo_url, tenant_living_from),
    )
    db.commit()
    return jsonify({"ok": True, "property_id": property_id})


@app.route("/api/flats")
def api_flats():
    db = get_db()
    block = (request.args.get("block") or "").strip()
    query = """
    SELECT p.id AS property_id, p.block, p.flat, COALESCE(od.owner_name, '') AS owner_name
    FROM properties p
    LEFT JOIN owner_details od ON od.property_id = p.id
    WHERE 1=1
    """
    params = []
    if block:
        query += " AND p.block = ?"
        params.append(block)
    query += " ORDER BY p.block, p.flat"
    rows = db.execute(query, params).fetchall()
    return jsonify({"ok": True, "data": [{
        "property_id": r["property_id"],
        "block": r["block"],
        "flat": r["flat"],
        "owner_name": r["owner_name"] or "",
    } for r in rows]})


@app.route("/api/tenants/<int:property_id>", methods=["PUT"])
def api_tenant_upsert(property_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    payload = request.get_json(silent=True) or {}
    exists = db.execute("SELECT id FROM properties WHERE id=?", (property_id,)).fetchone()
    if not exists:
        return jsonify({"ok": False, "error": "Flat not found"}), 404
    tenant_name = str(payload.get("tenant_name") or "").strip() or None
    tenant_contact = str(payload.get("tenant_contact") or "").strip() or None
    tenant_vehicle_list = str(payload.get("tenant_vehicle_list") or "").strip() or None
    tenant_photo_url = str(payload.get("tenant_photo_url") or "").strip() or None
    tenant_living_from = _normalize_date_for_storage(payload.get("tenant_living_from"))
    tenant_guard_payment_details = str(payload.get("tenant_guard_payment_details") or "").strip() or None
    db.execute(
        """
        INSERT INTO owner_details(property_id, is_occupied, occupied_by, tenant_name, tenant_contact, tenant_vehicle_list, tenant_photo_url, tenant_living_from, tenant_guard_payment_details, updated_at)
        VALUES (?, 1, 'TENANT', ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(property_id) DO UPDATE SET
          is_occupied=1,
          occupied_by='TENANT',
          tenant_name=excluded.tenant_name,
          tenant_contact=excluded.tenant_contact,
          tenant_vehicle_list=excluded.tenant_vehicle_list,
          tenant_photo_url=excluded.tenant_photo_url,
          tenant_living_from=excluded.tenant_living_from,
          tenant_guard_payment_details=excluded.tenant_guard_payment_details,
          updated_at=CURRENT_TIMESTAMP
        """,
        (property_id, tenant_name, tenant_contact, tenant_vehicle_list, tenant_photo_url, tenant_living_from, tenant_guard_payment_details),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/tenants/<int:property_id>", methods=["DELETE"])
def api_tenant_delete(property_id: int):
    admin = require_api_role("ADMIN")
    if admin:
        return admin
    db = get_db()
    db.execute(
        """
        UPDATE owner_details
        SET occupied_by='OWNER', tenant_name=NULL, tenant_contact=NULL, tenant_vehicle_list=NULL, tenant_photo_url=NULL, tenant_living_from=NULL, tenant_guard_payment_details=NULL, updated_at=CURRENT_TIMESTAMP
        WHERE property_id=?
        """,
        (property_id,),
    )
    db.commit()
    return jsonify({"ok": True})


if __name__ == "__main__":
    init_db()
    app.run(
        debug=(os.environ.get("FLASK_DEBUG", "false").lower() == "true"),
        host=os.environ.get("FLASK_HOST", "0.0.0.0"),
        port=int(os.environ.get("PORT", "5050")),
    )
